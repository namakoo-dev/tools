#!/usr/bin/env python
"""表計算を「見る」ための道具。描画して画像にする。

    python xlview.py book.xlsx              # 画像にしてパスを出す
    python xlview.py book.xlsx --dpi 150
    python xlview.py book.xlsx --sheet 月次  # 特定のシートだけ

## なぜ要るか

`openpyxl` で読めるのは値・数式・書式の**設定**であって、**結果の見た目**ではない。
見えないものが多い:

    グラフが正しく描けているか      ★ 全く見えない
    列幅が足りずに文字が切れているか  見えない
    条件付き書式の実際の色           規則は読めるが結果は見えない
    結合セル・印刷レイアウト         見えない

LibreOffice で PDF にして画像化すれば、全部見える。実測 (2026-08-06):
棒グラフ・軸・凡例・軸ラベル、条件付き書式の色階調、% 表示、
見出しの塗りと白文字 —— すべて読み取れた。

## ★ グラフが「無い」ように見えるとき

グラフシート上のグラフが出ないことがある。実測 (2026-08-06): openpyxl が作った
グラフシートは `<ext cx="0" cy="0"/>` で、**存在するが面積がゼロ**だった。
UNO で問い合わせるとグラフは 1 個ちゃんと読めているのに、描画すると白紙。

**「描画器が対応していない」と結論する前に、図形の大きさを見ること。**
(そう結論して 1 回外した。)

## ★ 既定のまま変換すると使いものにならない

印刷設定が既定だと、横に広い表は**列が複数ページに割れ**、余白の分だけ
**空白ページ**も出る (実測: 5 行 4 列の表で 2 ページ目が真っ白だった)。

だから変換前に「1 ページ幅に収める」を設定する。**★ 原本には触らない。
一時コピーに設定を書いてから変換する。**

## 前提

LibreOffice と pdftoppm (poppler)。
★ 専用プロファイルで起動する。既定プロファイルを使うと、GUI で開いている
LibreOffice を巻き込む。
"""
from __future__ import annotations

import argparse
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

PROFILE = Path(os.environ.get(
    "XLVIEW_PROFILE", str(Path.home() / ".nagi" / "lo-profile")))


def office_dir() -> Path:
    env = os.environ.get("XLVIEW_OFFICE")
    cands = [Path(env)] if env else []
    cands += [Path(r"C:\Program Files\LibreOffice\program"),
              Path(r"C:\Program Files (x86)\LibreOffice\program"),
              Path("/usr/lib/libreoffice/program"),
              Path("/Applications/LibreOffice.app/Contents/MacOS")]
    for c in cands:
        if (c / "soffice.exe").exists() or (c / "soffice").exists():
            return c
    raise SystemExit("LibreOffice が見つからない。XLVIEW_OFFICE で指定すること。")


def fit_to_page(src: Path, dst: Path, sheet: str | None) -> list[str]:
    """★ 原本に触らず、コピーに「1 ページ幅に収める」を設定する。

    戻り値は「こういう調整をした」の記録。見た目を変えた以上、黙らない。
    """
    import openpyxl

    notes: list[str] = []
    wb = openpyxl.load_workbook(src)
    targets = [wb[sheet]] if sheet else list(wb.worksheets)
    if sheet and sheet not in wb.sheetnames:
        raise SystemExit(f"シート {sheet!r} が無い。あるのは: {wb.sheetnames}")

    for ws in targets:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0      # 縦は成り行き。横だけ揃える
        ws.page_setup.orientation = "landscape"
        notes.append(f"{ws.title}: 横 1 ページに収めた (向きは横)")

    if sheet:
        for name in list(wb.sheetnames):
            if name != sheet:
                del wb[name]
        notes.append(f"{sheet} 以外のシートを外した")

    wb.save(dst)
    return notes


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("book")
    ap.add_argument("--sheet", default=None, help="このシートだけ見る")
    ap.add_argument("--dpi", type=int, default=120)
    ap.add_argument("--outdir", default=None, help="画像の置き場 (既定は一時領域)")
    ap.add_argument("--raw", action="store_true",
                    help="★ 印刷設定を触らずそのまま変換する (原本の見え方を確かめたいとき)")
    a = ap.parse_args()

    src = Path(a.book).resolve()
    if not src.exists():
        raise SystemExit(f"ファイルが無い: {src}")

    out = Path(a.outdir).resolve() if a.outdir else Path(tempfile.mkdtemp(prefix="xlview-"))
    out.mkdir(parents=True, exist_ok=True)

    work = out / ("view" + src.suffix)
    if a.raw:
        shutil.copy2(src, work)
        print("印刷設定は触っていない (--raw)", file=sys.stderr)
    else:
        for n in fit_to_page(src, work, a.sheet):
            # ★ 何を変えたかは必ず出す。黙って見た目を変えない
            print(f"調整: {n}", file=sys.stderr)

    PROFILE.mkdir(parents=True, exist_ok=True)
    soffice = office_dir() / ("soffice.exe" if os.name == "nt" else "soffice")
    proc = subprocess.run(
        [str(soffice), "--headless", "--norestore", "--convert-to", "pdf",
         "--outdir", str(out), str(work),
         f"-env:UserInstallation={PROFILE.resolve().as_uri()}"],
        capture_output=True, text=True)
    pdf = out / (work.stem + ".pdf")
    if not pdf.exists():
        sys.stderr.write(proc.stdout + proc.stderr)
        raise SystemExit("PDF にできなかった")

    for old in out.glob("page-*.jpg"):
        old.unlink()
    subprocess.run(["pdftoppm", "-jpeg", "-r", str(a.dpi), str(pdf),
                    str(out / "page")], check=True)

    pages = sorted(out.glob("page-*.jpg"))
    if not pages:
        raise SystemExit("画像にできなかった。pdftoppm は入っているか")
    print(f"{len(pages)} ページ")
    for p in pages:
        print(p)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
